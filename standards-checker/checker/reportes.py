import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .modelo import CUMPLE, NO_APLICA, NO_CUMPLE, ORDEN_SEVERIDAD, PARCIAL, REVISION

# La observacion va siempre al final de cada tabla.
COLUMNAS = ["PRS", "Micro", "Rama", "Criterio ID", "Seccion", "Criterio", "Resultado", "%", "Fuente",
            "Severidad", "Donde", "Observacion"]
COLOR = {CUMPLE: "C6EFCE", PARCIAL: "FFEB9C", NO_CUMPLE: "FFC7CE", REVISION: "DDEBF7", NO_APLICA: "EDEDED"}
SEV_COLOR = {"critica": "F4B6B6", "alta": "F8D7B0", "media": "FFF2CC", "baja": "E2EFDA"}


def _fila(prs, m, c, r):
    return [m.prs or prs, m.nombre, m.rama, c["id"], c["seccion"], c["descripcion"], r.estado,
            r.porcentaje if r.porcentaje is not None else "", "IA" if c["tipo"] == "ia" else "regla",
            c["severidad"], r.donde, r.obs]


def _filas_detalle(prs, micros):
    return [_fila(prs, m, c, r) for m in micros for c, r in m.resultados]


def _fmt(x):
    return "n/d" if x is None else f"{x:.1f}%"


def _opciones(c):
    return "\n".join(c.get("opciones", []))


def _donde(c, r):
    """Donde esta hoy el problema; si no hay archivos concretos, donde crear/ajustar lo que falta."""
    return r.donde or f"Crear/ajustar en: {c['donde_esperado']}"


def _pendientes(micros):
    """(micro, criterio, resultado) con incumplimiento, ordenados por severidad."""
    lista = [(m, c, r) for m in micros for c, r in m.resultados if r.estado in (NO_CUMPLE, PARCIAL)]
    lista.sort(key=lambda t: (ORDEN_SEVERIDAD[t[1]["severidad"]], t[0].nombre, t[1]["id"]))
    return lista


def generar_csv(prs, micros, ruta):
    with open(ruta, "w", newline="", encoding="utf-8-sig") as fh:  # utf-8-sig: Excel/Sheets abren bien los acentos
        w = csv.writer(fh)
        w.writerow(COLUMNAS)
        w.writerows(_filas_detalle(prs, micros))


def generar_json(prs, micros, ruta, meta):
    out = {"prs": prs, "meta": meta, "micros": []}
    for m in micros:
        out["micros"].append({
            "prs": m.prs or prs, "nombre": m.nombre, "rama": m.rama, "commit": m.commit, "error": m.error, "score": m.score(),
            "resultados": [{"id": c["id"], "estado": r.estado, "porcentaje": r.porcentaje, "obs": r.obs,
                            "donde": r.donde, "archivos": r.archivos} for c, r in m.resultados],
        })
    Path(ruta).write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")


def generar_excel(prs, micros, ruta, meta, criterios=None):
    wb = Workbook()
    negrita = Font(bold=True)
    cab = PatternFill("solid", fgColor="1F4E78")
    arriba = Alignment(wrap_text=True, vertical="top")

    def encabezado(ws, fila, textos):
        for i, t in enumerate(textos, 1):
            c = ws.cell(row=fila, column=i, value=t)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), cab
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    def anchos(ws, lista):
        for i, a in enumerate(lista, 1):
            ws.column_dimensions[get_column_letter(i)].width = a

    # ---- Resumen
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"], ws["A1"].font = "Verificacion del estandar de arquitectura", Font(bold=True, size=14)
    for i, (k, v) in enumerate([("PRS", prs), ("Fecha", meta["fecha"]), ("Estandar", meta["estandar"]),
                                ("Version de criterios", meta["version"])], start=2):
        ws.cell(row=i, column=1, value=k).font = negrita
        ws.cell(row=i, column=2, value=v)
    fila = 7
    encabezado(ws, fila, ["Micro", "Rama", "Commit", "Score", "Cumple", "Parcial", "No cumple",
                          "Revision manual", "No aplica", "Error"])
    for m in micros:
        fila += 1
        cu = m.conteo()
        for i, v in enumerate([m.nombre, m.rama, m.commit, m.score(), cu[CUMPLE], cu[PARCIAL], cu[NO_CUMPLE],
                               cu[REVISION], cu[NO_APLICA], m.error], 1):
            ws.cell(row=fila, column=i, value=v)
    ws.cell(row=fila + 2, column=1,
            value="Score = promedio ponderado por peso; excluye 'Requiere revision manual' y 'No aplica'.")
    anchos(ws, [28, 22, 10, 10, 9, 9, 10, 15, 10, 30])

    # ---- Detalle (observacion al final)
    wd = wb.create_sheet("Detalle")
    encabezado(wd, 1, COLUMNAS)
    filas = _filas_detalle(prs, micros)
    for r, fila_d in enumerate(filas, 2):
        for i, v in enumerate(fila_d, 1):
            wd.cell(row=r, column=i, value=v).alignment = arriba
        wd.cell(row=r, column=7).fill = PatternFill("solid", fgColor=COLOR.get(fila_d[6], "FFFFFF"))
        wd.cell(row=r, column=10).fill = PatternFill("solid", fgColor=SEV_COLOR.get(fila_d[9], "FFFFFF"))
    wd.freeze_panes = "A2"
    wd.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS))}{max(2, len(filas) + 1)}"
    anchos(wd, [8, 22, 20, 10, 24, 50, 14, 7, 8, 10, 55, 60])

    # ---- Plan de cambios: que cambiar, donde y como
    wp = wb.create_sheet("Plan de cambios")
    encabezado(wp, 1, ["Prioridad", "Micro", "Criterio ID", "Criterio", "Donde", "Que cambiar", "Opciones",
                       "Ejemplo", "Observacion"])
    for i, (m, c, r) in enumerate(_pendientes(micros), 2):
        for j, v in enumerate([c["severidad"], m.nombre, c["id"], c["descripcion"], _donde(c, r),
                               c["que_cambiar"], _opciones(c), c["ejemplo"], r.obs], 1):
            wp.cell(row=i, column=j, value=v).alignment = arriba
        wp.cell(row=i, column=1).fill = PatternFill("solid", fgColor=SEV_COLOR[c["severidad"]])
    wp.freeze_panes = "A2"
    anchos(wp, [10, 22, 10, 45, 50, 45, 70, 55, 55])

    # ---- Criterios: guia de correccion valida para cualquier micro
    if criterios:
        wc = wb.create_sheet("Criterios")
        encabezado(wc, 1, ["ID", "Seccion", "Criterio", "Como se verifica", "Que cambiar", "Donde debe quedar",
                           "Opciones", "Ejemplo", "Tipo", "Peso", "Severidad"])
        for i, c in enumerate(criterios, 2):
            for j, v in enumerate([c["id"], c["seccion"], c["descripcion"], c["verificacion"], c["que_cambiar"],
                                   c["donde_esperado"], _opciones(c), c["ejemplo"],
                                   "IA" if c["tipo"] == "ia" else "regla", c["peso"], c["severidad"]], 1):
                wc.cell(row=i, column=j, value=v).alignment = arriba
            wc.cell(row=i, column=11).fill = PatternFill("solid", fgColor=SEV_COLOR[c["severidad"]])
        wc.freeze_panes = "A2"
        anchos(wc, [9, 22, 45, 45, 45, 40, 70, 55, 7, 6, 11])
    wb.save(ruta)


def _esc(t):
    return str(t).replace("|", "\\|").replace("\n", " ")


def generar_md(prs, micros, ruta, meta):
    L = [f"# Reporte de verificacion del estandar de arquitectura — {prs}", "",
         f"- **Fecha:** {meta['fecha']}",
         f"- **Estandar:** {meta['estandar']} (criterios v{meta['version']})",
         "- **Alcance:** solo arquitectura. El score excluye los criterios *Requiere revision manual* y *No aplica*.",
         "- **Como usar este reporte:** cada hallazgo indica **donde** esta hoy el problema, **que cambiar** y "
         "**varias opciones** para hacerlo. La guia completa de cualquier criterio esta en `guia-correcciones.md`.",
         "", "## Resumen", "",
         "| Micro | Rama | Commit | Score | Cumple | Parcial | No cumple | Revision |",
         "|---|---|---|---|---|---|---|---|"]
    for m in micros:
        cu = m.conteo()
        L.append(f"| {m.nombre} | {m.rama} | {m.commit} | **{_fmt(m.score())}** | {cu[CUMPLE]} | {cu[PARCIAL]} | "
                 f"{cu[NO_CUMPLE]} | {cu[REVISION]} |")
    for m in micros:
        L += ["", f"## {m.nombre}" + (f" ({m.rama})" if m.rama else ""), ""]
        if m.error:
            L += [f"> **Requiere revision manual:** {_esc(m.error)}", ""]
            continue
        prio = [(c, r) for c, r in m.resultados if r.estado in (NO_CUMPLE, PARCIAL)]
        prio.sort(key=lambda t: (ORDEN_SEVERIDAD[t[0]["severidad"]], t[0]["id"]))
        if not prio:
            L += ["Sin hallazgos.", ""]
        for c, r in prio:
            L += [f"### `{c['id']}` · {c['severidad'].upper()} · {r.estado}", "",
                  f"**Criterio:** {c['descripcion']}", "",
                  f"- **Donde:** {_donde(c, r)}",
                  f"- **Que cambiar:** {c['que_cambiar']}",
                  "- **Opciones:**"]
            L += [f"  - {o}" for o in c["opciones"]]
            L += [f"- **Ejemplo:** `{c['ejemplo']}`" if len(c["ejemplo"]) < 160 else f"- **Ejemplo:** {c['ejemplo']}",
                  f"- **Observacion:** {r.obs}", ""]
        L += ["<details><summary>Todos los criterios</summary>", "",
              "| ID | Criterio | Sev. | Resultado | % | Observacion |", "|---|---|---|---|---|---|"]
        for c, r in m.resultados:
            pct = "" if r.porcentaje is None else f"{r.porcentaje:.0f}"
            L.append(f"| {c['id']} | {_esc(c['descripcion'])} | {c['severidad']} | {r.estado} | {pct} | {_esc(r.obs)} |")
        L += ["", "</details>"]
    Path(ruta).write_text("\n".join(L) + "\n", encoding="utf-8")


def generar_guia_md(criterios, ruta, meta):
    """Guia de correcciones generica (sirve para cualquier micro, aunque no se haya ejecutado el checker)."""
    L = ["# Guia de correcciones del estandar de arquitectura", "",
         f"Estandar: {meta['estandar']} (criterios v{meta['version']}). Para cada criterio: que cambiar, donde debe "
         "quedar, varias opciones y un ejemplo antes/despues. Sirve para cualquier micro de cualquier PRS.", "",
         "| ID | Criterio | Severidad |", "|---|---|---|"]
    L += [f"| [{c['id']}](#{c['id'].lower()}) | {_esc(c['descripcion'])} | {c['severidad']} |" for c in criterios]
    for c in criterios:
        L += ["", f"## {c['id']}", "", f"**{c['descripcion']}**", "",
              f"- **Severidad:** {c['severidad']} · **Peso:** {c['peso']} · **Se verifica:** "
              f"{'con IA / revision manual' if c['tipo'] == 'ia' else 'automaticamente'}",
              f"- **Como se verifica:** {c['verificacion']}",
              f"- **Que cambiar:** {c['que_cambiar']}",
              f"- **Donde debe quedar:** {c['donde_esperado']}", "- **Opciones:**"]
        L += [f"  - {o}" for o in c["opciones"]]
        L += [f"- **Ejemplo:** {c['ejemplo']}"]
    Path(ruta).write_text("\n".join(L) + "\n", encoding="utf-8")


def meta_ejecucion(version, estandar):
    return {"fecha": datetime.now().strftime("%Y-%m-%d %H:%M"), "version": version, "estandar": estandar}
